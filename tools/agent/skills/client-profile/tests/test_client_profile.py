"""Tests for the client-profile scripts. Run: python3 -m unittest discover -s tests (from the skill folder)."""
import copy
import datetime as dt
import json
import os
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(SKILL, "scripts"))
import validate_profile  # noqa: E402

with open(os.path.join(SKILL, "references", "example-profile.json"), encoding="utf-8") as f:
    EXAMPLE = json.load(f)
TODAY = dt.date(2026, 9, 29)


def write(path, data):
    with open(path, "w", encoding="utf-8") as f:
        f.write(data if isinstance(data, str) else json.dumps(data))


def read(path):
    with open(path, encoding="utf-8") as f:
        return f.read()


def check(profile):
    return validate_profile.validate(profile, today=TODAY)


def has(msgs, text):
    return any(text in m for m in msgs)


class Validator(unittest.TestCase):
    def test_example_is_clean(self):
        errors, warnings = check(copy.deepcopy(EXAMPLE))
        self.assertEqual(errors, [])
        self.assertEqual([w for w in warnings if "jsonschema" not in w], [])

    def test_percent_written_as_whole_number(self):
        p = copy.deepcopy(EXAMPLE); p["people"][0]["income"]["sgc_pct"] = 12
        errors, _ = check(p)
        self.assertTrue(has(errors, "/people/0/income/sgc_pct"))
        p = copy.deepcopy(EXAMPLE); p["accounts"][0]["performance"]["return_5yr_pa_pct"] = 7.1
        self.assertTrue(has(check(p)[0], "return_5yr_pa_pct"))

    def test_balance_without_provenance(self):
        p = copy.deepcopy(EXAMPLE); del p["provenance"]["/accounts/0/balance"]
        self.assertTrue(has(check(p)[0], "/accounts/0/balance: no provenance"))

    def test_an_override_counts_as_a_source(self):
        p = copy.deepcopy(EXAMPLE); del p["provenance"]["/accounts/0/balance"]
        p["overrides"] = [{"field": "/accounts/0/balance", "value": 275812.34, "who": "Sam Rivera", "date": "2026-09-01", "reason": "Confirmed by phone"}]
        self.assertEqual(check(p)[0], [])

    def test_provenance_cites_missing_document(self):
        p = copy.deepcopy(EXAMPLE); p["provenance"]["/accounts/0/balance"]["doc"] = "d99"
        self.assertTrue(has(check(p)[0], "cites 'd99'"))

    def test_insurer_recorded_as_super_fund(self):
        p = copy.deepcopy(EXAMPLE); p["accounts"][0]["provider"] = "TAL Insurance"
        self.assertTrue(has(check(p)[0], "is an insurer, not a super fund"))

    def test_mlc_super_is_a_fund(self):
        p = copy.deepcopy(EXAMPLE); p["accounts"][0]["provider"] = "MLC"; p["accounts"][0]["product"] = "MLC MasterKey Super Fundamentals"
        self.assertEqual(check(p)[0], [])

    def test_holdings_must_reconcile(self):
        p = copy.deepcopy(EXAMPLE); p["accounts"][1]["holdings"][1]["value"] = 50000; p["accounts"][1]["holdings"][1]["weight"] = 0.2
        warnings = check(p)[1]
        self.assertTrue(has(warnings, "weights sum to"))
        self.assertTrue(has(warnings, "values total"))

    def test_future_and_impossible_dates(self):
        p = copy.deepcopy(EXAMPLE); p["accounts"][0]["as_at"] = "2027-01-01"
        self.assertTrue(has(check(p)[0], "is in the future"))
        p = copy.deepcopy(EXAMPLE); p["people"][0]["dob"] = "1978-02-30"
        self.assertTrue(has(check(p)[0], "not a real date"))

    def test_nomination_expiry_may_be_in_the_future(self):
        self.assertIn("expires", EXAMPLE["accounts"][0]["beneficiary_nomination"])
        self.assertEqual(check(copy.deepcopy(EXAMPLE))[0], [])

    def test_owner_must_exist(self):
        p = copy.deepcopy(EXAMPLE); p["accounts"][1]["owner"] = "p9"
        self.assertTrue(has(check(p)[0], "'p9' is not a person id"))

    def test_defined_benefit_is_pointed_out(self):
        p = copy.deepcopy(EXAMPLE); p["accounts"][0]["account_type"] = "Defined Benefit"
        self.assertTrue(has(check(p)[1], "excluded from fee comparisons"))

    @unittest.skipUnless(__import__("importlib").util.find_spec("jsonschema"), "jsonschema not installed")
    def test_schema_catches_unknown_fields_and_bad_enums(self):
        p = copy.deepcopy(EXAMPLE); p["accounts"][0]["tfn"] = "123 456 789"
        self.assertTrue(has(check(p)[0], "Additional properties"))
        p = copy.deepcopy(EXAMPLE); p["risk_profile"]["profile"] = "Aggressive"
        self.assertTrue(has(check(p)[0], "/risk_profile/profile"))

    def test_cli_exit_codes(self):
        with tempfile.TemporaryDirectory() as d:
            good = os.path.join(d, "good.json"); bad = os.path.join(d, "bad.json")
            write(good, EXAMPLE)
            p = copy.deepcopy(EXAMPLE); p["accounts"][0]["provider"] = "Zurich"; write(bad, p)
            script = os.path.join(SKILL, "scripts", "validate_profile.py")
            self.assertEqual(subprocess.run([sys.executable, script, good], capture_output=True).returncode, 0)
            r = subprocess.run([sys.executable, script, bad], capture_output=True, text=True)
            self.assertEqual(r.returncode, 1); self.assertFalse(json.loads(r.stdout)["ok"])


class Extractor(unittest.TestCase):
    def run_extract(self, files):
        d = tempfile.mkdtemp()
        os.makedirs(os.path.join(d, "inputs", "documents"))
        for name, writer in files.items():
            writer(os.path.join(d, "inputs", "documents", name))
        r = subprocess.run([sys.executable, os.path.join(SKILL, "scripts", "extract.py")], cwd=d, capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stderr)
        inv = json.loads(r.stdout)
        return d, inv, {e["file"].split("/")[-1]: read(os.path.join(d, e["text"])) for e in inv}

    def test_csv_email_and_image(self):
        def csv_(p):
            write(p, "Code,Name,Units,Value\nVAS,Vanguard Australian Shares,600,59100\n")
        def eml(p):
            write(p, "From: Jane <jane@example.test>\nTo: adviser@example.test\nSubject: New salary\nDate: Tue, 1 Sep 2026 09:00:00 +1000\n\nHi, my salary is now $105,000.\n")
        def png(p):
            with open(p, "wb") as f:
                f.write(b"\x89PNG\r\n\x1a\n")
        _, inv, text = self.run_extract({"holdings.csv": csv_, "note.eml": eml, "portal.png": png})
        self.assertIn("A2=VAS", text["holdings.csv"]); self.assertIn("D2=59100", text["holdings.csv"])
        self.assertIn("Subject: New salary", text["note.eml"]); self.assertIn("$105,000", text["note.eml"])
        self.assertTrue(next(e for e in inv if e["file"].endswith("portal.png"))["needs_visual"])

    @unittest.skipUnless(__import__("importlib").util.find_spec("openpyxl"), "openpyxl not installed")
    def test_spreadsheet_cells_are_citable(self):
        def xlsx(p):
            import openpyxl
            wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Fact Finder"
            ws["A94"] = "Account balance"; ws["B94"] = 275812.34; ws["A15"] = "Date of birth"; ws["B15"] = dt.datetime(1978, 4, 12)
            wb.save(p)
        _, _, text = self.run_extract({"ff.xlsx": xlsx})
        self.assertIn("## Sheet: Fact Finder", text["ff.xlsx"])
        self.assertIn("A94=Account balance | B94=275812.34", text["ff.xlsx"])
        self.assertIn("B15=1978-04-12", text["ff.xlsx"])

    @unittest.skipUnless(__import__("importlib").util.find_spec("docx"), "python-docx not installed")
    def test_word_tables_are_citable(self):
        def docx_(p):
            import docx
            d = docx.Document(); d.add_paragraph("Client Goals Summary"); d.add_paragraph("Retire at 65")
            t = d.add_table(rows=2, cols=2); t.cell(0, 0).text = "Risk profile"; t.cell(0, 1).text = "High Growth"
            d.save(p)
        _, _, text = self.run_extract({"kyc.docx": docx_})
        self.assertIn("Retire at 65", text["kyc.docx"])
        self.assertIn("Table 1, row 1: Risk profile | High Growth", text["kyc.docx"])

    def test_one_bad_file_does_not_stop_the_rest(self):
        def broken(p):
            with open(p, "wb") as f:
                f.write(b"not really a spreadsheet")
        def txt(p):
            write(p, "hello")
        _, inv, text = self.run_extract({"broken.xlsx": broken, "note.txt": txt})
        self.assertEqual(next(e for e in inv if e["file"].endswith("broken.xlsx"))["note"], "extraction failed")
        self.assertIn("hello", text["note.txt"])


class Summary(unittest.TestCase):
    def test_summary_is_generated_from_the_json(self):
        with tempfile.TemporaryDirectory() as d:
            src = os.path.join(d, "p.json"); out = os.path.join(d, "p.md")
            write(src, EXAMPLE)
            subprocess.run([sys.executable, os.path.join(SKILL, "scripts", "summarise.py"), src, out], check=True, capture_output=True)
            md = read(out)
            self.assertIn("# Client profile: Jane & Tom Citizen", md)
            self.assertIn("$275,812", md)
            self.assertIn("**Total across accounts: $425,812**", md)
            self.assertIn("WARN", md)
            self.assertIn("Password-protected", md)


if __name__ == "__main__":
    unittest.main()
